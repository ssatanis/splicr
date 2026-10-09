"""Reproduce the published Millman CRISPRa count-table analysis in SplicR.

Run from the repository root with PYTHONPATH=engine engine/.tools/env/bin/python.
No database writes. All input and output files stay inside this study directory.
"""
import csv
import hashlib
import json
import os
import platform
import statistics
import subprocess
import time
from collections import Counter
from dataclasses import asdict
from datetime import datetime, timezone
from pathlib import Path

import openpyxl
from splicr.pipeline import ScreenInput, run_pipeline

ROOT = Path(__file__).resolve().parents[1]
started = datetime.now(timezone.utc).isoformat()
t_all = time.perf_counter()
source = ROOT / "source/Source_Data.xlsx"
wb = openpyxl.load_workbook(source, read_only=True, data_only=True)
sheet = wb["Whole-genome screen"]
header = next(sheet.iter_rows(values_only=True))
assert header[:7] == ('#seq.ID/s', 'sequence/s', 'D0', 'kidney', 'Gene Name', 'Promoter ID', 'Unique ID')
rows = []
seen = set()
for line, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
    gid, seq, d0, kidney, gene, promoter, uid = row[:7]
    assert gid and gid not in seen, (line, 'duplicate guide')
    assert seq and not (set(seq.upper()) - set('ACGT')), (line, 'invalid sequence')
    assert gene, (line, 'missing gene')
    assert all(isinstance(v, (int, float)) and v >= 0 and v == int(v) for v in [d0, kidney]), (line, 'invalid counts')
    seen.add(gid)
    rows.append(dict(guide_id=gid, sequence=seq.upper(), D0=int(d0), kidney=int(kidney), gene=gene,
                     is_control=gene == 'Non_Targeting_Human', promoter_id=promoter, original_unique_id=uid))
assert len(rows) == 106400
library = ROOT / 'input/study_library.tsv'
counts = ROOT / 'input/counts.tsv'
with library.open('w') as f, counts.open('w') as c:
    lw = csv.writer(f, delimiter='\t'); cw = csv.writer(c, delimiter='\t')
    lw.writerow(['guide_id', 'sequence', 'gene', 'is_control'])
    cw.writerow(['sgRNA', 'Gene', 'D0', 'kidney'])
    for r in rows:
        gene = '' if r['is_control'] else r['gene']
        lw.writerow([r['guide_id'], r['sequence'], gene, r['is_control']])
        cw.writerow([r['guide_id'], gene or 'CONTROL', r['D0'], r['kidney']])
library.with_suffix('.json').write_text(json.dumps({'name':'Cellecta KAHGW-106-P (published study guide map)', 'taxid':9606}))
gene_n = Counter(r['gene'] for r in rows if not r['is_control'])
audit = {'source_sha256':hashlib.sha256(source.read_bytes()).hexdigest(), 'guides':len(rows),
         'target_gene_labels':len(gene_n), 'non_targeting_guides':sum(r['is_control'] for r in rows),
         'guides_per_gene':dict(Counter(gene_n.values())), 'samples':{},
         'both_samples_zero':sum(r['D0']==0 and r['kidney']==0 for r in rows),
         'duplicate_sequences':len(rows)-len({r['sequence'] for r in rows}),
         'source_sheet':'Whole-genome screen', 'sample_limitation':'Only D0 and kidney aggregate columns are released. No mouse-level replicates or FASTQ files in this workbook.'}
for label in ['D0','kidney']:
    values = [r[label] for r in rows]
    audit['samples'][label] = {'total_counts':sum(values),'zero_guides':values.count(0),
        'detected_guides':len(values)-values.count(0), 'median_counts':statistics.median(values),
        'mean_counts':statistics.mean(values), 'guides_below_30_counts':sum(v<30 for v in values)}
authors = []
for row in wb['Fig.1'].iter_rows(min_row=3, values_only=True):
    if row[0] and isinstance(row[1],(int,float)):
        authors.append({'gene':row[0],'published_rra_score':row[1], 'published_plot_order':len(authors)+1})
with (ROOT/'input/published_enrichment_scores.csv').open('w') as f:
    dw=csv.DictWriter(f,fieldnames=list(authors[0])); dw.writeheader(); dw.writerows(authors)
wb.close()
(ROOT/'output/input_audit.json').write_text(json.dumps(audit,indent=2)+'\n')
conversion_seconds = time.perf_counter()-t_all
print('INPUT AUDIT', json.dumps(audit), flush=True)
spec = ScreenInput(name='Millman 2026 whole-genome CRISPRa kidney graft screen', count_table=counts,
    library_slug='millman_cellecta_kahgw106p', library_path=library,
    roles={'D0':'reference','kidney':'treatment'}, treatment=['kidney'], control=['D0'],
    normalization='median', fdr_threshold=0.1, hit_callers=['mageck_rra'],
    cell_line='HUES8-VPR stem-cell-derived islets', model_type='in_vivo',
    phenotype='transplantation enrichment', modality='crispra', condition='day 10 kidney capsule vs differentiated day 0', fitness_assay=False)
t_pipe = time.perf_counter()
result = run_pipeline(spec, ROOT/'output/splicr', persist=False, verbose=True)
pipeline_seconds = time.perf_counter()-t_pipe
finished = datetime.now(timezone.utc).isoformat()
timing={'started_at_utc':started,'finished_at_utc':finished,'input_conversion_and_audit_seconds':conversion_seconds,
        'splicr_pipeline_seconds':pipeline_seconds,'input_to_report_seconds':time.perf_counter()-t_all,
        'scope':'Workbook validation and conversion, then local SplicR count-table pipeline including QC, MAGeCK RRA, flags, atlas lookup and JSON report. Excludes download, dependency installation, standalone comparison, PDF creation and correspondence.',
        'machine':platform.platform(),'architecture':platform.machine(),'logical_cpus':os.cpu_count(),
        'processor':subprocess.run(['sysctl','-n','machdep.cpu.brand_string'],capture_output=True,text=True).stdout.strip(),
        'python':platform.python_version(), 'stages':[asdict(s) for s in result.stages],
        'ok':result.ok,'error':result.error,'summary':result.summary()}
(ROOT/'output/timing.json').write_text(json.dumps(timing,indent=2)+'\n')
print(result.summary(),flush=True)
print('MEASURED RUNTIME',json.dumps({k:timing[k] for k in ['input_conversion_and_audit_seconds','splicr_pipeline_seconds','input_to_report_seconds','ok']}),flush=True)
if not result.ok: raise SystemExit(1)
